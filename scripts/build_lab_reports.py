"""Build editable formal reports from the Markdown sources and retained cover."""
from pathlib import Path
import argparse,json,re
from copy import deepcopy
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.section import WD_SECTION,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/lab2'
OUT.mkdir(parents=True,exist_ok=True)
FIGURES=[('structure','Структурная схема системы'),('usecase','Диаграмма вариантов использования'),('activity_main','Основной сценарий'),('activity_alternative','Альтернативный сценарий замены блюда'),('activity_exception','Сценарий ошибки источника цен'),('wireframes','Предварительные экранные формы'),('er','ER-модель прототипа')]
FIGURES += [('entry_flow','Поток регистрации, профиля, дашборда и плана'),('onboarding_screens','Четыре экрана онбординга')]

FIGURES += [('er_domain','Предметная ER-модель и соответствие SQLite'),('activity_registration','Регистрация и заполнение профиля'),('activity_generation','Генерация меню')]

def clean(text):return re.sub(r'\*\*(.*?)\*\*',r'\1',text).replace('`','').replace('—','-')
def setfont(p,size=14,bold=False):
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
def configure(doc):
 for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(14);style.font.color.rgb=RGBColor(0,0,0)
  pf=style.paragraph_format;pf.line_spacing=1.5;pf.space_after=Pt(0);pf.space_before=Pt(0)
  if name.startswith('Heading'):
   style.font.bold=True;pf.space_before=Pt(10);pf.space_after=Pt(6);pf.keep_with_next=True
 doc.styles['Normal'].paragraph_format.widow_control=True
 sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(3);sec.right_margin=Cm(1.5);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.footer_distance=Cm(1)
 sec.different_first_page_header_footer=True
 for footer in [sec.footer,sec.first_page_footer]:
  for p in footer.paragraphs:p.text=''
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f);setfont(p,12)

def cover(lab,title):
 reference=Path('C:/Users/PC/Downloads/Лабораторная_2.docx')
 if reference.exists():
  doc=Document(reference);cut=doc.paragraphs[13]._p;remove=False
  for el in list(doc._element.body):
   if el is cut:remove=True
   if remove and el.tag!=qn('w:sectPr'):doc._element.body.remove(el)
  p=doc.paragraphs
  p[3].text='Кафедра «Программное обеспечение»';p[3].paragraph_format.space_after=Pt(75)
  p[4].text='ОТЧЁТ';p[4].style=doc.styles['Title'];p[4].alignment=WD_ALIGN_PARAGRAPH.CENTER
  p[5].text=f'по лабораторной работе № {lab}\n{title}'
  p[6].text='по дисциплине\n«Проектирование и конструирование ПО»\nПроект «Fit Nutrition Supply»';p[6].paragraph_format.space_after=Pt(35)
  p[7].text='';p[7].paragraph_format.space_after=Pt(0)
  p[12].text='Ижевск, 2026'
  for t in doc.tables:
   for row in t.rows:
    for c in row.cells:
     if 'Лугачев' in c.text:c.text='Еланцев М. О.'
     if 'старший преподаватель' in c.text:c.text='Проверил'
  for p in doc.paragraphs:
   setfont(p);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 else:
  doc=Document();p=doc.add_paragraph('МИНОБРНАУКИ РОССИИ\nИжевский государственный технический университет\nимени М. Т. Калашникова\nКафедра «Программное обеспечение»');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p=doc.add_paragraph(f'ОТЧЁТ\nпо лабораторной работе № {lab}\n{title}\nПроект «Fit Nutrition Supply»');p.paragraph_format.space_before=Pt(100);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  t=doc.add_table(rows=3,cols=2)
  for row,vals in zip(t.rows,[['Выполнили, Б23-191-2','Чипига А. А.\nФазлыев А. А.'],['Проверил','Еланцев М. О.'],['','Ижевск, 2026']]):
   for c,v in zip(row.cells,vals):c.text=v
 configure(doc)
 # Keep only relationships still used by the retained title (unrelated pictures are removed).
 used=set(re.findall(r'r:(?:embed|id)="([^"]+)"',doc._element.xml))
 for rid,rel in list(doc.part.rels.items()):
  if rel.reltype.endswith('/image') and rid not in used:doc.part.drop_rel(rid)
 doc.add_page_break();return doc

def paragraph(doc,text,small=False):
 p=doc.add_paragraph(clean(text));pf=p.paragraph_format;pf.first_line_indent=Cm(1.25);pf.line_spacing=1.5;p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if small:setfont(p,12);pf.line_spacing=1.15;pf.space_after=Pt(5)
 return p

def table(doc,rows,index):
 columns=len(rows[0]);widths={2:[4.5,12],3:[4.2,6.1,6.2],4:[3.8,4.2,4.3,4.2],5:[3.5,3.1,3.1,3.1,3.7]}[columns]
 p=doc.add_paragraph(f'Таблица {index}');setfont(p,12);p.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=0,cols=columns);t.style='Table Grid';t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for ri,row in enumerate(rows):
  cells=t.add_row().cells;pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for c,w,value in zip(cells,widths,row):
   c.width=Cm(w);c.text=clean(value)
   for p in c.paragraphs:
    p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.first_line_indent=Pt(0);setfont(p,10.5 if columns==5 else 11,ri==0)
    if ri==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.add_paragraph().paragraph_format.space_after=Pt(3)

def parse(doc,text,adr=False,counters=None,compact=False):
 counters=counters or {'tables':0};lines=text.splitlines();i=0;code=False
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line.startswith('```'):code=not code;continue
  if code:
   p=doc.add_paragraph(line);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(2)
   p.paragraph_format.keep_with_next=i<len(lines) and not lines[i].strip().startswith('```')
   for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
   continue
  if line.startswith('[[diagram:'):continue
  if line=='[[adr:all]]':
   for f in sorted((ROOT/'docs/lab2/adr').glob('*.md')):
    counters['adrs']=counters.get('adrs',0)+1;parse(doc,f.read_text(encoding='utf-8'),True,counters)
   continue
  if line.startswith('Рисунок '):
   paragraph(doc,'См. '+line[0].lower()+line[1:]+' (приложение А).',True);continue
  if line.startswith('|'):
   raw=[line]
   while i<len(lines) and lines[i].strip().startswith('|'):raw.append(lines[i].strip());i+=1
   rows=[[cell.strip() for cell in r.strip('|').split('|')] for r in raw if not re.fullmatch(r'[| :\-]+',r)]
   counters['tables']+=1;table(doc,rows,counters['tables']);continue
  if line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));title=clean(line.lstrip('# '))
   if level==1 and not adr:continue
   style=('Heading 2' if level==1 else 'Heading 3') if adr else 'Heading 2' if compact else 'Heading 1' if level<=2 else 'Heading 2' if level==3 else 'Heading 3'
   p=doc.add_paragraph(title,style=style)
   if style=='Heading 1':p.paragraph_format.page_break_before=True
   if adr:
    setfont(p,14 if level==1 else 12,True);p.paragraph_format.line_spacing=1.2
    p.paragraph_format.space_before=Pt(8 if level==1 else 4);p.paragraph_format.space_after=Pt(4 if level==1 else 2)
    if level==1:p.paragraph_format.page_break_before=counters.get('adrs',0)>1
   continue
  p=paragraph(doc,line,small=bool(re.match(r'^\d+\. ',line) and 'http' in line))
  if line.endswith('JSON-запрос:') or line.endswith('проекция:') or line.endswith('плана:') or line.endswith('снимка:') or line.endswith('записи:'):p.paragraph_format.keep_with_next=True
  if adr or compact:setfont(p,12);p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(2)

def append_figures(doc):
 for index,(name,title) in enumerate(FIGURES,1):
  sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.different_first_page_header_footer=False
  wide=name in ('usecase','wireframes','er','entry_flow','onboarding_screens','er_domain')
  landscape=wide or name=='structure'
  if name in ('activity_alternative','activity_registration','activity_generation'):sec.page_width=Cm(29.7);sec.page_height=Cm(42)
  elif wide:sec.page_width=Cm(42);sec.page_height=Cm(29.7)
  elif landscape:sec.page_width=Cm(29.7);sec.page_height=Cm(21)
  else:sec.page_width=Cm(21);sec.page_height=Cm(29.7)
  sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
  sec.left_margin=Cm(2);sec.right_margin=Cm(1.5);sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5)
  p=doc.add_paragraph('Приложение А Схемы' if index==1 else 'Приложение А Продолжение');setfont(p,12);p.paragraph_format.space_after=Pt(12)
  image_path=ROOT/'docs/lab2/diagrams'/f'{name}.png'
  with Image.open(image_path) as im:iw,ih=im.size
  maxw=int(sec.page_width-sec.left_margin-sec.right_margin);maxh=int(sec.page_height-sec.top_margin-sec.bottom_margin-Cm(3))
  width=min(maxw,int(maxh*iw/ih))
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
  p.add_run().add_picture(str(image_path),width=width)
  p=doc.add_paragraph(f'Рисунок {index} - {title}');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;setfont(p,12);p.paragraph_format.line_spacing=1.15

def build():
 doc=cover(2,'Архитектура и интерактивный прототип в коде')
 p=doc.add_paragraph('СОДЕРЖАНИЕ',style='Heading 1');p.paragraph_format.page_break_before=False
 entries=['Введение','1 Уточнение концепции после замечаний','2 Структурная схема и технологии','3 Сценарии использования и интерфейс','4 Модель данных','5 Программные интерфейсы','6 Архитектурные решения','7 Проверка прототипа и покрытие требований','8 Работа с ИИ агентом','Заключение','Источники','Приложение А Схемы','Приложение Б Протокол проверки']
 pagination=ROOT/'docs/lab2/report-pagination.json'
 mapping=json.loads(pagination.read_text(encoding='utf-8')) if pagination.exists() else {}
 for title in entries:
  p=doc.add_paragraph();p.paragraph_format.line_spacing=1.5;p.paragraph_format.tab_stops.add_tab_stop(Cm(16.4),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  p.add_run(title+'\t'+str(mapping.get(title,'')));setfont(p,12)
 parse(doc,(ROOT/'docs/lab2/architecture.md').read_text(encoding='utf-8'))
 append_figures(doc)
 # Return to A4 after the diagrams and keep the complete test evidence editable.
 sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.orientation=WD_ORIENT.PORTRAIT
 sec.left_margin=Cm(3);sec.right_margin=Cm(1.5);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
 p=doc.add_paragraph('Приложение Б Протокол проверки',style='Heading 1');p.paragraph_format.page_break_before=False
 text=(ROOT/'docs/lab2/test_results.md').read_text(encoding='utf-8');parse(doc,text,compact=True)
 doc.save(OUT/'Fit_Nutrition_Supply_LR2.docx')
 doc=cover(1,'Дополнение к отчёту\nКонцепция и спецификация проекта')
 parse(doc,(ROOT/'docs/lab1_revision.md').read_text(encoding='utf-8'))
 doc.save(OUT/'Fit_Nutrition_Supply_LR1_Corrections.docx')
 print(OUT)
if __name__=='__main__':build()
